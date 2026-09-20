#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""مولد docx عربي RTL بنسق المرجع الزراعي — وثيقة 44"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Cairo'
INK = RGBColor(0x1A, 0x1D, 0x2E)
BRAND = RGBColor(0x7C, 0x5C, 0xE6)

def rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def set_font(run, size=12, bold=False, color=INK):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:cs'), FONT); rFonts.set(qn('w:ascii'), FONT); rFonts.set(qn('w:hAnsi'), FONT)
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'), str(size*2)); rPr.append(cs)

def h(doc, text, size=18, color=BRAND):
    p = doc.add_paragraph(); rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run(text), size, True, color); return p

def body(doc, text, size=12):
    p = doc.add_paragraph(); rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run(text), size); return p

def table(doc, headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # RTL للجدول كله
    tblPr = t._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual'); tblPr.append(bidi)
    for j, hd in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ''
        p = c.paragraphs[0]; rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(hd), 11, True, BRAND)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i+1].cells[j]; c.text = ''
            p = c.paragraphs[0]; rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_font(p.add_run(str(val)), 11)
    return t

if __name__ == '__main__':
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
    # صفحة العينة — بوابة و٠
    h(doc, 'عائلتي — Family OS', 26)
    h(doc, 'صفحة عينة الأنماط (بوابة و٠)', 16, INK)
    h(doc, '١-١ عنوان فرعي من المستوى الثاني', 14)
    body(doc, 'هذه فقرة نصية عادية بخط Cairo واتجاه من اليمين إلى اليسار، تحاكي أسلوب فقرات المرجع الزراعي. يعد نظام «عائلتي» منصة تربوية متكاملة تربط إنجاز الأبناء بعملة الدقائق، وتمنح الأب تحكمًا كاملًا وشفافًا بلا خداع.')
    body(doc, 'وتتفرع المشكلة العامة إلى المشكلات الآتية: (عينة من نمط التفريع في المرجع).')
    h(doc, 'جدول ١-١: عينة بنسق «المرحلة/الأعمال/المخرج المتوقع»', 12, INK)
    table(doc,
        ['المرحلة', 'الأعمال التي ستُنفَّذ', 'المخرج المتوقع'],
        [['التحليل', 'جرد الخدمات والرحلات من السجل الرسمي', '٢٤٠ خدمة موزعة على ٥ دومينات'],
         ['النمذجة', 'مخططات UML: حالات استخدام وتسلسل ونشاط', 'حزمة drawio + SVG'],
         ['التصميم', 'ERD وجداول الحقل/الوصف من schema.sql', '٢٠ جدول قاعدة بيانات موثقًا']])
    doc.add_paragraph()
    h(doc, 'جدول ١-٢: عينة بنسق «الحقل/الوصف»', 12, INK)
    table(doc, ['الحقل', 'الوصف'],
        [['id', 'معرف فريد للسجل (uuid)'],
         ['family_id', 'مفتاح أجنبي إلى جدول العائلة'],
         ['role', 'دور العضو: OWNER أو PARENT أو GUARDIAN']])
    doc.save('صفحة_العينة_و0.docx')
    print('✅ صفحة_العينة_و0.docx')
